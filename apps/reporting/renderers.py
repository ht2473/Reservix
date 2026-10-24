"""Писатели форматов выгрузки."""

from __future__ import annotations

import csv
import io
from functools import lru_cache
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from apps.reporting.structure import (
    Column,
    ReportDocument,
    Table,
    ValueKind,
    as_number,
    format_value,
)
from core.logging import get_logger

logger = get_logger("reporting")


class RenderError(Exception):
    """Отчёт не удалось записать в запрошенном формате."""


#: Имя, под которым кириллический шрифт регистрируется в reportlab.
PDF_FONT_NAME = "ReservixSans"
PDF_FONT_BOLD_NAME = "ReservixSans-Bold"


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def render_docx(document: ReportDocument) -> bytes:
    """Записывает отчёт в формат .docx."""
    from docx import Document as DocxDocument
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    docx = DocxDocument()

    widest = max((len(table.columns) for table in document.tables), default=0)
    if widest > 6:
        section = docx.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    style = docx.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    if document.organization:
        header = docx.add_paragraph(document.organization)
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    docx.add_heading(document.title, level=1)
    if document.subtitle:
        docx.add_paragraph(document.subtitle)

    _docx_requisites(docx, document)

    if document.metrics:
        docx.add_heading("Ключевые показатели", level=2)
        for metric in document.metrics:
            line = docx.add_paragraph(style="List Bullet")
            line.add_run(f"{metric.label}: ").bold = True
            line.add_run(format_value(metric.value, metric.kind))
            if metric.note:
                line.add_run(f" — {metric.note}")

    for table in document.tables:
        docx.add_heading(table.title, level=2)
        if table.note:
            docx.add_paragraph(table.note)
        _docx_table(docx, table)

    if document.notes:
        docx.add_heading("Примечания", level=2)
        for note in document.notes:
            docx.add_paragraph(note, style="List Bullet")

    _docx_signature(docx, document)

    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


def _docx_requisites(docx, document: ReportDocument) -> None:
    """Выводит дату формирования и параметры отбора."""
    moment = document.generated_at or timezone.localtime()
    lines = [f"Сформирован: {moment.strftime('%d.%m.%Y %H:%M')}"]
    lines.extend(f"{label}: {value}" for label, value in document.parameters)

    paragraph = docx.add_paragraph()
    run = paragraph.add_run("\n".join(lines))
    run.italic = True


def _docx_table(docx, table: Table) -> None:
    """Записывает таблицу с шапкой и полосами сетки."""
    if not table.rows:
        docx.add_paragraph("Данных, отвечающих условиям отбора, не найдено.")
        return

    grid = docx.add_table(rows=1, cols=len(table.columns))
    grid.style = "Light Grid Accent 1"

    for cell, column in zip(grid.rows[0].cells, table.columns, strict=True):
        cell.text = column.title
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row in table.rows:
        cells = grid.add_row().cells
        for cell, column in zip(cells, table.columns, strict=True):
            cell.text = format_value(row.get(column.key), column.kind)


def _docx_signature(docx, document: ReportDocument) -> None:
    """Добавляет строку подписи: отчёт кадровой службы подписывается."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    docx.add_paragraph()
    line = docx.add_paragraph(f"Отчёт сформирован: {document.author or '—'}")
    line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signature = docx.add_paragraph("Подпись ____________________")
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
#: Форматы ячеек по типу столбца. Числовой формат задаёт вид, сохраняя тип
#: значения: столбец остаётся пригодным для сортировки и сводных таблиц.
XLSX_NUMBER_FORMATS: dict[ValueKind, str] = {
    ValueKind.INTEGER: "# ##0",
    ValueKind.NUMBER: "# ##0.00",
    ValueKind.PERCENT: '0.0" %"',
    ValueKind.DATE: "DD.MM.YYYY",
}


def render_xlsx(document: ReportDocument) -> bytes:
    """Записывает отчёт в формат .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="EEF2FF")

    summary["A1"] = document.title
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = document.subtitle
    summary["A3"] = document.organization

    moment = document.generated_at or timezone.localtime()
    summary["A5"] = "Сформирован"
    summary["B5"] = moment.strftime("%d.%m.%Y %H:%M")
    summary["A6"] = "Автор"
    summary["B6"] = document.author

    line = 8
    if document.parameters:
        summary.cell(row=line, column=1, value="Параметры отбора").font = bold
        line += 1
        for label, value in document.parameters:
            summary.cell(row=line, column=1, value=label)
            summary.cell(row=line, column=2, value=value)
            line += 1
        line += 1

    if document.metrics:
        summary.cell(row=line, column=1, value="Ключевые показатели").font = bold
        line += 1
        for metric in document.metrics:
            summary.cell(row=line, column=1, value=metric.label)
            cell = summary.cell(row=line, column=2)
            _xlsx_write(cell, metric.value, metric.kind)
            summary.cell(row=line, column=3, value=metric.note)
            line += 1

    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["C"].width = 48

    for index, table in enumerate(document.tables, start=1):
        sheet = workbook.create_sheet(_sheet_title(table.title, index))

        for position, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=1, column=position, value=column.title)
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            width = column.width or max(len(column.title) + 4, 14)
            sheet.column_dimensions[get_column_letter(position)].width = width

        for row_index, row in enumerate(table.rows, start=2):
            for position, column in enumerate(table.columns, start=1):
                _xlsx_write(
                    sheet.cell(row=row_index, column=position),
                    row.get(column.key),
                    column.kind,
                )

        sheet.freeze_panes = "A2"
        if table.rows:
            last = get_column_letter(len(table.columns))
            sheet.auto_filter.ref = f"A1:{last}{len(table.rows) + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_write(cell, value, kind: ValueKind) -> None:
    """Записывает значение в ячейку, сохраняя его тип."""
    number_format = XLSX_NUMBER_FORMATS.get(kind)

    if kind in {ValueKind.NUMBER, ValueKind.INTEGER, ValueKind.PERCENT}:
        number = as_number(value)
        cell.value = number
        if number is not None and number_format:
            cell.number_format = number_format
        return

    if kind is ValueKind.DATE and value is not None:
        cell.value = value
        cell.number_format = XLSX_NUMBER_FORMATS[ValueKind.DATE]
        return

    cell.value = format_value(value, kind) if value is not None else None


def _sheet_title(title: str, index: int) -> str:
    """Приводит заголовок к допустимому имени листа."""
    cleaned = "".join(" " if symbol in "[]:*?/\\" else symbol for symbol in title).strip()
    cleaned = cleaned[:28] or "Лист"
    return f"{index}. {cleaned}"[:31]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
@lru_cache(maxsize=1)
def resolve_pdf_font() -> str:
    """Регистрирует кириллический шрифт и возвращает его имя."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if PDF_FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return PDF_FONT_NAME

    candidates = [Path(path) for path in settings.PDF_FONT_PATHS if path]
    for path in candidates:
        if not path.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(path)))
        except Exception as error:  # pragma: no cover - повреждённый файл шрифта
            logger.warning("Шрифт не подошёл", extra={"path": str(path), "error": str(error)})
            continue

        bold = _bold_variant(path)
        if bold is not None:
            pdfmetrics.registerFont(TTFont(PDF_FONT_BOLD_NAME, str(bold)))
        else:
            # Полужирного начертания рядом нет: заголовки печатаются обычным.
            pdfmetrics.registerFont(TTFont(PDF_FONT_BOLD_NAME, str(path)))

        logger.info("Шрифт PDF зарегистрирован", extra={"path": str(path)})
        return PDF_FONT_NAME

    raise RenderError(
        "Не найден шрифт с поддержкой кириллицы для выгрузки в PDF. "
        "Укажите путь к файлу .ttf в переменной окружения PDF_FONT_PATH."
    )


def pdf_font_available() -> bool:
    """Доступна ли выгрузка в PDF в текущем окружении."""
    try:
        resolve_pdf_font()
    except RenderError:
        return False
    return True


def _bold_variant(regular: Path) -> Path | None:
    """Ищет полужирное начертание рядом с обычным по принятым суффиксам."""
    stem = regular.stem
    for suffix in ("-Bold", "bd", "-bold", "Bold"):
        candidate = regular.with_name(f"{stem}{suffix}{regular.suffix}")
        if candidate.exists():
            return candidate
    return None


def render_pdf(document: ReportDocument) -> bytes:
    """Записывает отчёт в формат PDF."""
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    font = resolve_pdf_font()
    widest = max((len(table.columns) for table in document.tables), default=0)
    page_size = landscape(A4) if widest > 6 else A4

    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=document.title,
        author=document.author,
    )

    sample = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=sample["Normal"], fontName=font, fontSize=9, leading=12)
    head = ParagraphStyle(
        "Head", parent=sample["Heading1"], fontName=PDF_FONT_BOLD_NAME, fontSize=15, leading=19
    )
    sub = ParagraphStyle(
        "Sub", parent=sample["Heading2"], fontName=PDF_FONT_BOLD_NAME, fontSize=11, leading=15
    )
    right = ParagraphStyle("Right", parent=body, alignment=TA_RIGHT)

    story: list = []

    if document.organization:
        story.append(Paragraph(document.organization, right))
    story.append(Paragraph(document.title, head))
    if document.subtitle:
        story.append(Paragraph(document.subtitle, body))

    moment = document.generated_at or timezone.localtime()
    story.append(Paragraph(f"Сформирован: {moment.strftime('%d.%m.%Y %H:%M')}", body))
    for label, value in document.parameters:
        story.append(Paragraph(f"{label}: {value}", body))
    story.append(Spacer(1, 6 * mm))

    if document.metrics:
        story.append(Paragraph("Ключевые показатели", sub))
        for metric in document.metrics:
            text = f"<b>{metric.label}</b>: {format_value(metric.value, metric.kind)}"
            if metric.note:
                text += f" — {metric.note}"
            story.append(Paragraph(text, body))
        story.append(Spacer(1, 5 * mm))

    available = page_size[0] - pdf.leftMargin - pdf.rightMargin

    for index, table in enumerate(document.tables):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(table.title, sub))
        if table.note:
            story.append(Paragraph(table.note, body))
        story.append(Spacer(1, 3 * mm))

        if not table.rows:
            story.append(Paragraph("Данных, отвечающих условиям отбора, не найдено.", body))
            continue

        story.append(_pdf_table(table, body, available, font))

    if document.notes:
        story.append(Spacer(1, 5 * mm))
        story.append(Paragraph("Примечания", sub))
        for note in document.notes:
            story.append(Paragraph(note, body))

    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph(f"Отчёт сформирован: {document.author or '—'}", right))

    pdf.build(story)
    return buffer.getvalue()


def _pdf_table(table: Table, body_style, available: float, font: str):
    """Собирает таблицу PDF с переносом текста внутри ячеек."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, TableStyle
    from reportlab.platypus import Table as PdfTable

    cell_style = ParagraphStyle("Cell", parent=body_style, fontSize=8, leading=10)
    header_style = ParagraphStyle(
        "CellHead", parent=cell_style, fontName=PDF_FONT_BOLD_NAME, textColor=colors.white
    )

    header = [Paragraph(column.title, header_style) for column in table.columns]
    data = [header]
    for row in table.rows:
        data.append(
            [
                Paragraph(format_value(row.get(column.key), column.kind), cell_style)
                for column in table.columns
            ]
        )

    widths = _column_widths(table.columns, available)
    grid = PdfTable(data, colWidths=widths, repeatRows=1)
    grid.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312E81")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7CBD9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return grid


def _column_widths(columns: list[Column], available: float) -> list[float]:
    """Распределяет ширину страницы между столбцами пропорционально объявленной."""
    declared = [column.width or 12 for column in columns]
    total = sum(declared)
    return [available * value / total for value in declared]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def render_csv(document: ReportDocument) -> bytes:
    """Записывает первую таблицу отчёта в CSV."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

    if not document.tables:
        writer.writerow([document.title])
        return buffer.getvalue().encode("utf-8-sig")

    table = document.tables[0]
    writer.writerow([column.title for column in table.columns])
    for row in table.rows:
        writer.writerow(
            [format_value(row.get(column.key), column.kind) for column in table.columns]
        )

    return buffer.getvalue().encode("utf-8-sig")


#: Соответствие формата выгрузки и писателя.
RENDERERS = {
    "docx": render_docx,
    "xlsx": render_xlsx,
    "pdf": render_pdf,
    "csv": render_csv,
}

#: Тип содержимого для отдачи файла браузеру.
CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv; charset=utf-8",
}


def render(document: ReportDocument, output_format: str) -> bytes:
    """Записывает отчёт в указанном формате."""
    writer = RENDERERS.get(output_format)
    if writer is None:
        raise RenderError(f"Формат выгрузки не поддерживается: {output_format}")
    return writer(document)
