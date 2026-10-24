"""Модульные тесты нейтрального описания отчёта и писателей форматов."""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import pytest

from apps.reporting.renderers import (
    RenderError,
    pdf_font_available,
    render,
    render_csv,
    render_docx,
    render_pdf,
    render_xlsx,
)
from apps.reporting.structure import (
    Column,
    Metric,
    ReportDocument,
    Table,
    ValueKind,
    as_number,
    format_value,
)

pytestmark = pytest.mark.unit


@pytest.fixture
def document() -> ReportDocument:
    """Отчёт со всеми типами столбцов и непустой сводкой."""
    report = ReportDocument(
        title="Состав кадрового резерва",
        subtitle="Проверочная выгрузка",
        organization="ООО «Меридиан Софт»",
        author="Иванов И. И.",
        generated_at=dt.datetime(2026, 3, 14, 9, 30),
        parameters=[("Период с", "01.01.2026"), ("Группа резерва", "управленческий")],
        metrics=[
            Metric("Состоит в резерве", 42, ValueKind.INTEGER),
            Metric("Доля готовых", 27.5, ValueKind.PERCENT, "от состава резерва"),
        ],
    )
    report.add_table(
        Table(
            title="Резервисты",
            columns=[
                Column("employee", "Резервист", width=24),
                Column("index", "Индекс", ValueKind.NUMBER, width=10),
                Column("share", "Обеспечено весом", ValueKind.PERCENT, width=12),
                Column("included", "Включён", ValueKind.DATE, width=12),
                Column("ready", "Готов", ValueKind.BOOLEAN, width=8),
            ],
            rows=[
                {
                    "employee": "Смирнов Игорь Владимирович",
                    "index": Decimal("78.40"),
                    "share": 83.3,
                    "included": dt.date(2025, 6, 1),
                    "ready": True,
                },
                {
                    "employee": "Волкова Анна Петровна",
                    "index": None,
                    "share": None,
                    "included": None,
                    "ready": False,
                },
            ],
            note="Проверочные данные.",
        )
    )
    report.notes.append("Отчёт сформирован для проверки писателей форматов.")
    return report


# ---------------------------------------------------------------------------
# Оформление значений
# ---------------------------------------------------------------------------
class TestFormatValue:
    """Приведение значений к виду, принятому в русскоязычном документе."""

    def test_целое_разделяет_разряды_неразрывным_пробелом(self):
        """
        Разделитель — именно неразрывный пробел (U+00A0): в напечатанном документе разряды числа не
        должны переноситься на разные строки.
        """
        assert format_value(1240567, ValueKind.INTEGER) == "1 240 567"

    def test_дробное_использует_запятую(self):
        assert format_value(Decimal("78.4"), ValueKind.NUMBER) == "78,40"

    def test_процент_получает_знак(self):
        assert format_value(27.53, ValueKind.PERCENT) == "27,5 %"

    def test_дата_в_российском_формате(self):
        assert format_value(dt.date(2026, 3, 14), ValueKind.DATE) == "14.03.2026"

    def test_пустое_значение_становится_прочерком(self):
        """
        Пустая ячейка в напечатанном отчёте неотличима от пропущенной при заполнении, и проверяющий
        обязан различать «сведений нет» и «строка не дописана».
        """
        assert format_value(None, ValueKind.NUMBER) == "—"
        assert format_value("", ValueKind.TEXT) == "—"

    def test_логическое_читается_словом(self):
        assert format_value(True, ValueKind.BOOLEAN) == "да"
        assert format_value(False, ValueKind.BOOLEAN) == "нет"


class TestAsNumber:
    """Приведение значения к числу для ячейки таблицы Excel."""

    def test_дробная_строка_с_запятой_разбирается(self):
        assert as_number("78,4") == pytest.approx(78.4)

    def test_нечисловое_значение_даёт_пустоту(self):
        """
        Строка в числовом столбце превращает его в текстовый для всей колонки, поэтому
        неразобранное значение не записывается вовсе.
        """
        assert as_number("готов") is None
        assert as_number(None) is None


class TestDocumentStructure:
    """Свойства описания отчёта."""

    def test_число_строк_считается_по_всем_таблицам(self, document: ReportDocument):
        document.add_table(
            Table(title="Вторая", columns=[Column("a", "А")], rows=[{"a": 1}, {"a": 2}])
        )
        assert document.total_rows == 4

    def test_значения_строки_следуют_порядку_столбцов(self, document: ReportDocument):
        table = document.tables[0]
        values = table.values(table.rows[0])
        assert values[0] == "Смирнов Игорь Владимирович"
        assert values[4] is True


# ---------------------------------------------------------------------------
# Писатели форматов
# ---------------------------------------------------------------------------
class TestDocxWriter:
    """Выгрузка в формат .docx."""

    def test_файл_читается_библиотекой_обратно(self, document: ReportDocument):
        from docx import Document as DocxDocument

        content = render_docx(document)
        assert content[:2] == b"PK"  # контейнер OOXML — архив ZIP

        parsed = DocxDocument(io.BytesIO(content))
        texts = [paragraph.text for paragraph in parsed.paragraphs]
        assert document.title in texts
        assert any("Меридиан" in text for text in texts)

    def test_таблица_содержит_шапку_и_все_строки(self, document: ReportDocument):
        from docx import Document as DocxDocument

        parsed = DocxDocument(io.BytesIO(render_docx(document)))
        grid = parsed.tables[0]
        assert grid.rows[0].cells[0].text == "Резервист"
        assert len(grid.rows) == 3  # шапка и две строки данных

    def test_пустая_таблица_объясняет_отсутствие_данных(self):
        """
        Пустой раздел без пояснения читается как сбой формирования, а не как отсутствие подходящих
        записей.
        """
        from docx import Document as DocxDocument

        report = ReportDocument(title="Пустой отчёт")
        report.add_table(Table(title="Резервисты", columns=[Column("a", "А")]))

        parsed = DocxDocument(io.BytesIO(render_docx(report)))
        assert any("не найдено" in paragraph.text for paragraph in parsed.paragraphs)


class TestXlsxWriter:
    """Выгрузка в формат .xlsx."""

    def test_числовые_столбцы_записаны_числами(self, document: ReportDocument):
        """
        Записанное строкой число не суммируется и не сортируется — выгрузка перестаёт быть рабочим
        материалом и становится картинкой таблицы.
        """
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(render_xlsx(document)))
        sheet = book["1. Резервисты"]

        assert sheet["B2"].value == pytest.approx(78.4)
        assert isinstance(sheet["B2"].value, int | float)
        assert sheet["C2"].value == pytest.approx(83.3)

    def test_первый_лист_содержит_реквизиты_и_показатели(self, document: ReportDocument):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(render_xlsx(document)))
        summary = book["Сводка"]
        assert summary["A1"].value == document.title
        values = [cell.value for row in summary.iter_rows() for cell in row]
        assert "Состоит в резерве" in values

    def test_шапка_закреплена_и_снабжена_автофильтром(self, document: ReportDocument):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(render_xlsx(document)))
        sheet = book["1. Резервисты"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:E3"

    def test_пустое_значение_не_превращает_столбец_в_текстовый(self, document: ReportDocument):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(render_xlsx(document)))
        sheet = book["1. Резервисты"]
        assert sheet["B3"].value is None


class TestCsvWriter:
    """Выгрузка в формат CSV."""

    def test_кодировка_и_разделитель_понятны_табличному_редактору(self, document: ReportDocument):
        """
        Файл с запятыми и без метки порядка байтов открывается в русской локали одной колонкой с
        искажёнными буквами, и пользователь считает выгрузку неисправной.
        """
        content = render_csv(document)
        assert content.startswith(b"\xef\xbb\xbf")

        text = content.decode("utf-8-sig")
        assert text.splitlines()[0] == "Резервист;Индекс;Обеспечено весом;Включён;Готов"

    def test_записывается_первая_таблица(self, document: ReportDocument):
        text = render_csv(document).decode("utf-8-sig")
        assert "Смирнов Игорь Владимирович" in text
        assert len(text.strip().splitlines()) == 3


@pytest.mark.skipif(
    not pdf_font_available(),
    reason="в окружении нет шрифта с поддержкой кириллицы",
)
class TestPdfWriter:
    """Выгрузка в формат PDF."""

    def test_файл_имеет_признак_формата(self, document: ReportDocument):
        content = render_pdf(document)
        assert content.startswith(b"%PDF-")
        assert len(content) > 1000

    def test_широкая_таблица_печатается_на_альбомной_странице(self):
        """
        Отчёт с девятью столбцами на книжной странице нечитаем; проверяется сам факт формирования —
        размер страницы задаётся до сборки документа.
        """
        report = ReportDocument(title="Широкий отчёт")
        report.add_table(
            Table(
                title="Много столбцов",
                columns=[Column(f"c{i}", f"Столбец {i}") for i in range(9)],
                rows=[{f"c{i}": i for i in range(9)}],
            )
        )
        assert render_pdf(report).startswith(b"%PDF-")


class TestRenderDispatch:
    """Выбор писателя по формату."""

    def test_известный_формат_обслуживается(self, document: ReportDocument):
        assert render(document, "csv").startswith(b"\xef\xbb\xbf")

    def test_неизвестный_формат_отвергается(self, document: ReportDocument):
        with pytest.raises(RenderError):
            render(document, "rtf")
